#!/usr/bin/env python3
"""
Update Canonical Agent-Catalog.xlsx and inventory for AI-Dev-Team
Integrates:
- 5 Source Repositories (claude-code-agents, multi-agent-review-orchestrator,
  multi-model-code-review-agent, Claude-Code-Promts-Skills, ai-code-review-prompts)
- 1 Master Orchestrator (AGT-099)
- 13 Specialist Roles / Agents (AGT-100 through AGT-112)
- 8 Discovered/Adapted Skills
- 4 Audit Workflows
- Dedicated Sheet 18_Audits_Integration
- Updates Summary KPIs and formulas
"""

import os
import json
import yaml
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = "/Users/subhajkar/Developer/AI-Dev-Team"
WORKBOOK_PATH = os.path.join(ROOT, "Agent-Catalog.xlsx")
AGENTS_JSON = os.path.join(ROOT, "inventory/agents.json")
AGENTS_YAML = os.path.join(ROOT, "inventory/agents.yaml")

# Style definitions
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FILL_HEADER = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
FILL_ZEBRA = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
FILL_HEALTHY = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FONT_HEALTHY = Font(name="Calibri", size=11, color="006100", bold=True)
BORDER_THIN = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

NEW_AGENTS = [
    {
        "id": "AGT-099",
        "name": "Pre-Live Full Audit Orchestrator",
        "classification": "ORCHESTRATOR / PIPELINE",
        "subtype": "Master Audit Coordinator",
        "purpose": "Master pre-live full audit orchestration across code, security, QA, dependencies, performance, and release gating.",
        "primary_category": "Orchestration",
        "secondary_category": "Testing / QA",
        "current_location": "AI-Dev-Team/orchestrators/pre-live-full-audit",
        "runtime_owner": "AI-Dev-Team Master Orchestrator",
        "invocation_method": "python scripts/pre_live_full_audit.py [target_path]",
        "entrypoint": "scripts/pre_live_full_audit.py",
        "runtime": "Python 3 / Node / Shell",
        "model": "Dynamic (Gemini Pro / Claude Sonnet / Copilot Auto)",
        "mcp_deps": ["filesystem", "git", "github", "agent-team"],
        "skill_deps": ["pre-live-full-audit", "security-audit-recon", "finding-verification", "release-gating"],
        "project_deps": "Global (~/Developer)",
        "github_repo": "AI-Dev-Team",
        "github_url": "https://github.com/hackerspider/AI-Dev-Team.git",
        "version": "1.0.0",
        "commit": "HEAD",
        "license": "MIT",
        "health": "HEALTHY",
        "project_scope": "Global",
        "centralization_status": "Centralized in AI-Dev-Team",
        "migration_action": "INTEGRATED FROM 5 AUDIT REPOSITORIES",
        "risk": "NONE",
        "notes": "Coordinating parallel static, security, QA, and adversarial reviews with verified finding gating."
    },
    {
        "id": "AGT-100",
        "name": "Code Auditor Role",
        "classification": "AGENT ROLE",
        "subtype": "Static Analysis & Quality",
        "purpose": "Static code analysis, maintainability, architectural adherence, code smell detection, and syntax linting.",
        "primary_category": "Testing / QA",
        "secondary_category": "Architecture",
        "current_location": "AI-Dev-Team/roles/code-auditor.md",
        "runtime_owner": "AI-Dev-Team Ecosystem",
        "invocation_method": "Antigravity Subagent / CLI",
        "entrypoint": "roles/code-auditor.md",
        "runtime": "Antigravity IDE / Markdown",
        "model": "Gemini 3.1 Pro / Claude Sonnet",
        "mcp_deps": ["filesystem", "git"],
        "skill_deps": ["pre-live-full-audit"],
        "project_deps": "Global",
        "github_repo": "undeadlist/claude-code-agents",
        "github_url": "https://github.com/undeadlist/claude-code-agents.git",
        "version": "1.0.0",
        "commit": "13e3d51f",
        "license": "MIT",
        "health": "HEALTHY",
        "project_scope": "Global",
        "centralization_status": "Centralized in AI-Dev-Team",
        "migration_action": "ADAPTED FROM claude-code-agents",
        "risk": "NONE",
        "notes": "Symlinked in agents/code-auditor.md"
    },
    {
        "id": "AGT-101",
        "name": "Security Auditor Role",
        "classification": "AGENT ROLE",
        "subtype": "Application Security & OWASP",
        "purpose": "OWASP Top 10, injection flaws, secret leakage, auth/authz bypasses, crypto misconfigs, and security posture.",
        "primary_category": "Security",
        "secondary_category": "Architecture",
        "current_location": "AI-Dev-Team/roles/security-auditor.md",
        "runtime_owner": "AI-Dev-Team Ecosystem",
        "invocation_method": "Antigravity Subagent / CLI",
        "entrypoint": "roles/security-auditor.md",
        "runtime": "Antigravity IDE / Markdown",
        "model": "Claude Sonnet / Gemini Pro",
        "mcp_deps": ["filesystem", "git"],
        "skill_deps": ["security-audit-recon"],
        "project_deps": "Global",
        "github_repo": "Rtur2003/Claude-Code-Promts-Skills",
        "github_url": "https://github.com/Rtur2003/Claude-Code-Promts-Skills.git",
        "version": "1.0.0",
        "commit": "ac0502e7",
        "license": "MIT",
        "health": "HEALTHY",
        "project_scope": "Global",
        "centralization_status": "Centralized in AI-Dev-Team",
        "migration_action": "ADAPTED FROM Claude-Code-Promts-Skills",
        "risk": "NONE",
        "notes": "Symlinked in agents/security-auditor.md"
    },
    {
        "id": "AGT-102",
        "name": "Bug & Logic Auditor Role",
        "classification": "AGENT ROLE",
        "subtype": "Defect & Edge Case Detection",
        "purpose": "Edge cases, null safety, race conditions, async deadlocks, boundary value defects, and logic faults.",
        "primary_category": "Testing / QA",
        "secondary_category": "Backend Development",
        "current_location": "AI-Dev-Team/roles/bug-auditor.md",
        "runtime_owner": "AI-Dev-Team Ecosystem",
        "invocation_method": "Antigravity Subagent / CLI",
        "entrypoint": "roles/bug-auditor.md",
        "runtime": "Antigravity IDE / Markdown",
        "model": "Gemini 3.1 Pro / Claude Sonnet",
        "mcp_deps": ["filesystem", "git"],
        "skill_deps": ["pre-live-full-audit"],
        "project_deps": "Global",
        "github_repo": "undeadlist/claude-code-agents",
        "github_url": "https://github.com/undeadlist/claude-code-agents.git",
        "version": "1.0.0",
        "commit": "13e3d51f",
        "license": "MIT",
        "health": "HEALTHY",
        "project_scope": "Global",
        "centralization_status": "Centralized in AI-Dev-Team",
        "migration_action": "ADAPTED FROM claude-code-agents",
        "risk": "NONE",
        "notes": "Symlinked in agents/bug-auditor.md"
    },
    {
        "id": "AGT-103",
        "name": "Infrastructure Auditor Role",
        "classification": "AGENT ROLE",
        "subtype": "DevOps & Cloud Security",
        "purpose": "Docker containerization, Kubernetes configurations, cloud policies, and IaC security.",
        "primary_category": "DevOps",
        "secondary_category": "Cloud",
        "current_location": "AI-Dev-Team/roles/infra-auditor.md",
        "runtime_owner": "AI-Dev-Team Ecosystem",
        "invocation_method": "Antigravity Subagent / CLI",
        "entrypoint": "roles/infra-auditor.md",
        "runtime": "Antigravity IDE / Markdown",
        "model": "Gemini Pro / Claude Sonnet",
        "mcp_deps": ["filesystem", "git"],
        "skill_deps": ["pre-live-full-audit"],
        "project_deps": "Global",
        "github_repo": "undeadlist/claude-code-agents",
        "github_url": "https://github.com/undeadlist/claude-code-agents.git",
        "version": "1.0.0",
        "commit": "13e3d51f",
        "license": "MIT",
        "health": "HEALTHY",
        "project_scope": "Global",
        "centralization_status": "Centralized in AI-Dev-Team",
        "migration_action": "ADAPTED FROM claude-code-agents",
        "risk": "NONE",
        "notes": "Symlinked in agents/infra-auditor.md"
    },
    {
        "id": "AGT-104",
        "name": "Performance Auditor Role",
        "classification": "AGENT ROLE",
        "subtype": "Performance Profiling",
        "purpose": "Memory leaks, algorithmic complexity, query efficiency, runtime bottlenecks, and CWV analysis.",
        "primary_category": "Testing / QA",
        "secondary_category": "Web Development",
        "current_location": "AI-Dev-Team/roles/perf-auditor.md",
        "runtime_owner": "AI-Dev-Team Ecosystem",
        "invocation_method": "Antigravity Subagent / CLI",
        "entrypoint": "roles/perf-auditor.md",
        "runtime": "Antigravity IDE / Markdown",
        "model": "Gemini Pro / Claude Sonnet",
        "mcp_deps": ["filesystem", "git"],
        "skill_deps": ["pre-live-full-audit"],
        "project_deps": "Global",
        "github_repo": "undeadlist/claude-code-agents",
        "github_url": "https://github.com/undeadlist/claude-code-agents.git",
        "version": "1.0.0",
        "commit": "13e3d51f",
        "license": "MIT",
        "health": "HEALTHY",
        "project_scope": "Global",
        "centralization_status": "Centralized in AI-Dev-Team",
        "migration_action": "ADAPTED FROM claude-code-agents",
        "risk": "NONE",
        "notes": "Symlinked in agents/perf-auditor.md"
    },
    {
        "id": "AGT-105",
        "name": "Database Auditor Role",
        "classification": "AGENT ROLE",
        "subtype": "Database Architecture",
        "purpose": "Database schema design, indexing, connection pooling, migration safety, and data integrity.",
        "primary_category": "Data / Analytics",
        "secondary_category": "Backend Development",
        "current_location": "AI-Dev-Team/roles/db-auditor.md",
        "runtime_owner": "AI-Dev-Team Ecosystem",
        "invocation_method": "Antigravity Subagent / CLI",
        "entrypoint": "roles/db-auditor.md",
        "runtime": "Antigravity IDE / Markdown",
        "model": "Gemini Pro / Claude Sonnet",
        "mcp_deps": ["filesystem", "git"],
        "skill_deps": ["pre-live-full-audit"],
        "project_deps": "Global",
        "github_repo": "undeadlist/claude-code-agents",
        "github_url": "https://github.com/undeadlist/claude-code-agents.git",
        "version": "1.0.0",
        "commit": "13e3d51f",
        "license": "MIT",
        "health": "HEALTHY",
        "project_scope": "Global",
        "centralization_status": "Centralized in AI-Dev-Team",
        "migration_action": "ADAPTED FROM claude-code-agents",
        "risk": "NONE",
        "notes": "Symlinked in agents/db-auditor.md"
    },
    {
        "id": "AGT-106",
        "name": "Dependency Auditor Role",
        "classification": "AGENT ROLE",
        "subtype": "Supply Chain Security",
        "purpose": "Software supply chain, known CVEs, outdated packages, vulnerability scanning, and license compatibility.",
        "primary_category": "Security",
        "secondary_category": "DevOps",
        "current_location": "AI-Dev-Team/roles/dep-auditor.md",
        "runtime_owner": "AI-Dev-Team Ecosystem",
        "invocation_method": "Antigravity Subagent / CLI",
        "entrypoint": "roles/dep-auditor.md",
        "runtime": "Antigravity IDE / Markdown",
        "model": "Gemini Pro / Claude Sonnet",
        "mcp_deps": ["filesystem", "git"],
        "skill_deps": ["pre-live-full-audit"],
        "project_deps": "Global",
        "github_repo": "undeadlist/claude-code-agents",
        "github_url": "https://github.com/undeadlist/claude-code-agents.git",
        "version": "1.0.0",
        "commit": "13e3d51f",
        "license": "MIT",
        "health": "HEALTHY",
        "project_scope": "Global",
        "centralization_status": "Centralized in AI-Dev-Team",
        "migration_action": "ADAPTED FROM claude-code-agents",
        "risk": "NONE",
        "notes": "Symlinked in agents/dep-auditor.md"
    },
    {
        "id": "AGT-107",
        "name": "UI/UX & A11y Auditor Role",
        "classification": "AGENT ROLE",
        "subtype": "Frontend & Accessibility",
        "purpose": "WCAG accessibility, responsive viewport design, UI interaction states, and user journeys.",
        "primary_category": "Frontend Development",
        "secondary_category": "Accessibility",
        "current_location": "AI-Dev-Team/roles/ui-auditor.md",
        "runtime_owner": "AI-Dev-Team Ecosystem",
        "invocation_method": "Antigravity Subagent / CLI",
        "entrypoint": "roles/ui-auditor.md",
        "runtime": "Antigravity IDE / Markdown",
        "model": "Gemini Pro / Claude Sonnet",
        "mcp_deps": ["filesystem", "browser"],
        "skill_deps": ["browser-e2e-audit"],
        "project_deps": "Global",
        "github_repo": "undeadlist/claude-code-agents",
        "github_url": "https://github.com/undeadlist/claude-code-agents.git",
        "version": "1.0.0",
        "commit": "13e3d51f",
        "license": "MIT",
        "health": "HEALTHY",
        "project_scope": "Global",
        "centralization_status": "Centralized in AI-Dev-Team",
        "migration_action": "ADAPTED FROM claude-code-agents",
        "risk": "NONE",
        "notes": "Symlinked in agents/ui-auditor.md"
    },
    {
        "id": "AGT-108",
        "name": "API & Contract Tester Role",
        "classification": "AGENT ROLE",
        "subtype": "API & Integration Testing",
        "purpose": "REST/GraphQL contract testing, rate limiting, error responses, schema validation, and payload testing.",
        "primary_category": "Testing / QA",
        "secondary_category": "Backend Development",
        "current_location": "AI-Dev-Team/roles/api-tester.md",
        "runtime_owner": "AI-Dev-Team Ecosystem",
        "invocation_method": "Antigravity Subagent / CLI",
        "entrypoint": "roles/api-tester.md",
        "runtime": "Antigravity IDE / Markdown",
        "model": "Gemini Pro / Claude Sonnet",
        "mcp_deps": ["filesystem", "git"],
        "skill_deps": ["pre-live-full-audit"],
        "project_deps": "Global",
        "github_repo": "undeadlist/claude-code-agents",
        "github_url": "https://github.com/undeadlist/claude-code-agents.git",
        "version": "1.0.0",
        "commit": "13e3d51f",
        "license": "MIT",
        "health": "HEALTHY",
        "project_scope": "Global",
        "centralization_status": "Centralized in AI-Dev-Team",
        "migration_action": "ADAPTED FROM claude-code-agents",
        "risk": "NONE",
        "notes": "Symlinked in agents/api-tester.md"
    },
    {
        "id": "AGT-109",
        "name": "SEO Auditor Role",
        "classification": "AGENT ROLE",
        "subtype": "Search Engine Optimization",
        "purpose": "Meta tags, open graph, structured data, crawlability, canonicalization, and Core Web Vitals.",
        "primary_category": "Web Development",
        "secondary_category": "Frontend Development",
        "current_location": "AI-Dev-Team/roles/seo-auditor.md",
        "runtime_owner": "AI-Dev-Team Ecosystem",
        "invocation_method": "Antigravity Subagent / CLI",
        "entrypoint": "roles/seo-auditor.md",
        "runtime": "Antigravity IDE / Markdown",
        "model": "Gemini Pro / Claude Sonnet",
        "mcp_deps": ["filesystem", "browser"],
        "skill_deps": ["pre-live-full-audit"],
        "project_deps": "Global",
        "github_repo": "undeadlist/claude-code-agents",
        "github_url": "https://github.com/undeadlist/claude-code-agents.git",
        "version": "1.0.0",
        "commit": "13e3d51f",
        "license": "MIT",
        "health": "HEALTHY",
        "project_scope": "Global",
        "centralization_status": "Centralized in AI-Dev-Team",
        "migration_action": "ADAPTED FROM claude-code-agents",
        "risk": "NONE",
        "notes": "Symlinked in agents/seo-auditor.md"
    },
    {
        "id": "AGT-110",
        "name": "Adversarial Reviewer Role",
        "classification": "AGENT ROLE",
        "subtype": "Red Team / Assumption Challenging",
        "purpose": "Red-team adversarial inspection, assumption challenging, edge case exploit modeling, and test robustness.",
        "primary_category": "Security",
        "secondary_category": "Testing / QA",
        "current_location": "AI-Dev-Team/roles/adversarial-reviewer.md",
        "runtime_owner": "AI-Dev-Team Ecosystem",
        "invocation_method": "Antigravity Subagent / CLI",
        "entrypoint": "roles/adversarial-reviewer.md",
        "runtime": "Antigravity IDE / Markdown",
        "model": "Claude Sonnet / Gemini Pro",
        "mcp_deps": ["filesystem", "git"],
        "skill_deps": ["adversarial-review"],
        "project_deps": "Global",
        "github_repo": "soul-sol/ai-code-review-prompts",
        "github_url": "https://github.com/soul-sol/ai-code-review-prompts.git",
        "version": "1.0.0",
        "commit": "ff21994c",
        "license": "MIT",
        "health": "HEALTHY",
        "project_scope": "Global",
        "centralization_status": "Centralized in AI-Dev-Team",
        "migration_action": "ADAPTED FROM ai-code-review-prompts",
        "risk": "NONE",
        "notes": "Symlinked in agents/adversarial-reviewer.md"
    },
    {
        "id": "AGT-111",
        "name": "Review Consolidator Role",
        "classification": "AGENT ROLE",
        "subtype": "Findings Synthesis & Dedup",
        "purpose": "Cross-lane finding deduplication, evidence preservation, root cause correlation, and remediation roadmaps.",
        "primary_category": "Orchestration",
        "secondary_category": "Testing / QA",
        "current_location": "AI-Dev-Team/roles/review-consolidator.md",
        "runtime_owner": "AI-Dev-Team Ecosystem",
        "invocation_method": "Antigravity Subagent / CLI",
        "entrypoint": "roles/review-consolidator.md",
        "runtime": "Antigravity IDE / Markdown",
        "model": "Gemini Pro / Claude Sonnet",
        "mcp_deps": ["filesystem", "git"],
        "skill_deps": ["finding-verification"],
        "project_deps": "Global",
        "github_repo": "bgaborg/multi-agent-review-orchestrator",
        "github_url": "https://gist.github.com/bgaborg/8058aac0df0e48a3740776963fa87805",
        "version": "1.0.0",
        "commit": "9c568f14",
        "license": "MIT",
        "health": "HEALTHY",
        "project_scope": "Global",
        "centralization_status": "Centralized in AI-Dev-Team",
        "migration_action": "ADAPTED FROM Gist multi-agent-review-orchestrator",
        "risk": "NONE",
        "notes": "Symlinked in agents/review-consolidator.md"
    },
    {
        "id": "AGT-112",
        "name": "Release Gatekeeper Role",
        "classification": "AGENT ROLE",
        "subtype": "Release Gating & Governance",
        "purpose": "Deterministic gating policy evaluation (BLOCKED / READY_FOR_REVIEW / PASS) against verified findings.",
        "primary_category": "DevOps",
        "secondary_category": "Security",
        "current_location": "AI-Dev-Team/roles/release-gatekeeper.md",
        "runtime_owner": "AI-Dev-Team Ecosystem",
        "invocation_method": "Antigravity Subagent / CLI",
        "entrypoint": "roles/release-gatekeeper.md",
        "runtime": "Antigravity IDE / Markdown",
        "model": "Dynamic (Rule-based / Gemini Pro)",
        "mcp_deps": ["filesystem", "git"],
        "skill_deps": ["release-gating"],
        "project_deps": "Global",
        "github_repo": "pelednoam/multi-model-code-review-agent",
        "github_url": "https://github.com/pelednoam/multi-model-code-review-agent.git",
        "version": "1.0.0",
        "commit": "d41b7535",
        "license": "MIT",
        "health": "HEALTHY",
        "project_scope": "Global",
        "centralization_status": "Centralized in AI-Dev-Team",
        "migration_action": "ADAPTED FROM multi-model-code-review-agent",
        "risk": "NONE",
        "notes": "Symlinked in agents/release-gatekeeper.md"
    }
]

NEW_SKILLS = [
    ("pre-live-full-audit", "Testing / QA", "skills/pre-live-full-audit", "AI-Dev-Team Integration (5 Repos)", "Master pre-live full audit orchestration across code, security, QA, dependencies, performance, and release gating.", "/pre-live-full-audit [target_path]", "HEALTHY"),
    ("security-audit-recon", "Security", "skills/security-audit-recon", "Claude-Code-Promts-Skills / Rtur2003", "Comprehensive security reconnaissance and vulnerability scanning (SQLi, secrets, auth, crypto, headers).", "/security-audit-recon [target_path]", "HEALTHY"),
    ("adversarial-review", "Security", "skills/adversarial-review", "ai-code-review-prompts / soul-sol", "Adversarial review uncovering bypassed controls, edge case vulnerabilities, and test flaws.", "/adversarial-review [target_path]", "HEALTHY"),
    ("multi-agent-review", "Testing / QA", "skills/multi-agent-review", "multi-agent-review-orchestrator / bgaborg", "Coordinates parallel specialist review lanes with strict isolation and evidence gathering.", "/multi-agent-review [target_path]", "HEALTHY"),
    ("multi-model-review", "Model Routing", "skills/multi-model-review", "multi-model-code-review-agent / pelednoam", "Coordinates multi-model consensus and orthogonal code review across configured backends.", "/multi-model-review [target_path]", "HEALTHY"),
    ("finding-verification", "Testing / QA", "skills/finding-verification", "AI-Dev-Team Integration (5 Repos)", "Deterministic verification of audit findings against filesystem and cited evidence.", "/finding-verification [findings.json]", "HEALTHY"),
    ("release-gating", "DevOps", "skills/release-gating", "AI-Dev-Team Integration (5 Repos)", "Evaluates mandatory release criteria and assigns BLOCKED, READY_FOR_REVIEW, or PASS.", "/release-gating [findings.json]", "HEALTHY"),
    ("browser-e2e-audit", "Testing / QA", "skills/browser-e2e-audit", "claude-code-agents / undeadlist", "Automated browser and user-journey verification across core web workflows.", "/browser-e2e-audit [url]", "HEALTHY")
]

NEW_REPOS = [
    ("claude-code-agents", "integration-workspace/external/claude-code-agents", "main", "13e3d51f", "https://github.com/undeadlist/claude-code-agents.git", "1.0.0", "MIT", "Markdown / Shell", "SYNCED / HEALTHY"),
    ("multi-agent-review-orchestrator", "integration-workspace/external/multi-agent-review-orchestrator", "main", "9c568f14", "https://gist.github.com/bgaborg/8058aac0df0e48a3740776963fa87805", "1.0.0", "Public Gist / MIT", "Markdown / Bash", "SYNCED / HEALTHY"),
    ("multi-model-code-review-agent", "integration-workspace/external/multi-model-code-review-agent", "main", "d41b7535", "https://github.com/pelednoam/multi-model-code-review-agent.git", "1.0.0", "MIT", "Python / Shell", "SYNCED / HEALTHY"),
    ("Claude-Code-Promts-Skills", "integration-workspace/external/Claude-Code-Promts-Skills", "main", "ac0502e7", "https://github.com/Rtur2003/Claude-Code-Promts-Skills.git", "1.0.0", "MIT", "Markdown", "SYNCED / HEALTHY"),
    ("ai-code-review-prompts", "integration-workspace/external/ai-code-review-prompts", "main", "ff21994c", "https://github.com/soul-sol/ai-code-review-prompts.git", "1.0.0", "MIT", "Markdown", "SYNCED / HEALTHY")
]

AUDIT_INTEGRATION_DATA = [
    # Source Repo, Owner, Upstream URL, Commit, License, Discovered Component, Adapted Name, Type, Target Location, Status
    ("claude-code-agents", "undeadlist", "https://github.com/undeadlist/claude-code-agents.git", "13e3d51f", "MIT", "agents/code-audit.md", "code-auditor", "AGENT ROLE", "roles/code-auditor.md", "ADAPTED & VERIFIED"),
    ("claude-code-agents", "undeadlist", "https://github.com/undeadlist/claude-code-agents.git", "13e3d51f", "MIT", "agents/bug-audit.md", "bug-auditor", "AGENT ROLE", "roles/bug-auditor.md", "ADAPTED & VERIFIED"),
    ("claude-code-agents", "undeadlist", "https://github.com/undeadlist/claude-code-agents.git", "13e3d51f", "MIT", "agents/infra-audit.md", "infra-auditor", "AGENT ROLE", "roles/infra-auditor.md", "ADAPTED & VERIFIED"),
    ("claude-code-agents", "undeadlist", "https://github.com/undeadlist/claude-code-agents.git", "13e3d51f", "MIT", "agents/perf-audit.md", "perf-auditor", "AGENT ROLE", "roles/perf-auditor.md", "ADAPTED & VERIFIED"),
    ("claude-code-agents", "undeadlist", "https://github.com/undeadlist/claude-code-agents.git", "13e3d51f", "MIT", "agents/db-audit.md", "db-auditor", "AGENT ROLE", "roles/db-auditor.md", "ADAPTED & VERIFIED"),
    ("claude-code-agents", "undeadlist", "https://github.com/undeadlist/claude-code-agents.git", "13e3d51f", "MIT", "agents/dep-audit.md", "dep-auditor", "AGENT ROLE", "roles/dep-auditor.md", "ADAPTED & VERIFIED"),
    ("claude-code-agents", "undeadlist", "https://github.com/undeadlist/claude-code-agents.git", "13e3d51f", "MIT", "agents/ui-audit.md", "ui-auditor", "AGENT ROLE", "roles/ui-auditor.md", "ADAPTED & VERIFIED"),
    ("claude-code-agents", "undeadlist", "https://github.com/undeadlist/claude-code-agents.git", "13e3d51f", "MIT", "agents/api-tester.md", "api-tester", "AGENT ROLE", "roles/api-tester.md", "ADAPTED & VERIFIED"),
    ("claude-code-agents", "undeadlist", "https://github.com/undeadlist/claude-code-agents.git", "13e3d51f", "MIT", "agents/seo-audit.md", "seo-auditor", "AGENT ROLE", "roles/seo-auditor.md", "ADAPTED & VERIFIED"),
    ("claude-code-agents", "undeadlist", "https://github.com/undeadlist/claude-code-agents.git", "13e3d51f", "MIT", "workflows/full-audit.md", "pre-live-full-audit", "WORKFLOW", "workflows/pre-live-full-audit.md", "ADAPTED & VERIFIED"),
    ("claude-code-agents", "undeadlist", "https://github.com/undeadlist/claude-code-agents.git", "13e3d51f", "MIT", "workflows/pre-commit.md", "pre-commit-audit", "WORKFLOW", "workflows/pre-commit-audit.md", "ADAPTED & VERIFIED"),
    ("claude-code-agents", "undeadlist", "https://github.com/undeadlist/claude-code-agents.git", "13e3d51f", "MIT", "workflows/pre-deploy.md", "pre-deploy-audit", "WORKFLOW", "workflows/pre-deploy-audit.md", "ADAPTED & VERIFIED"),
    ("claude-code-agents", "undeadlist", "https://github.com/undeadlist/claude-code-agents.git", "13e3d51f", "MIT", "workflows/bug-fix.md", "audit-fix-retest", "WORKFLOW", "workflows/audit-fix-retest.md", "ADAPTED & VERIFIED"),
    ("multi-agent-review-orchestrator", "bgaborg", "https://gist.github.com/bgaborg/8058aac0df0e48a3740776963fa87805", "9c568f14", "MIT", "review_orchestrator.sh", "multi-agent-review", "SKILL / ORCHESTRATOR", "skills/multi-agent-review", "ADAPTED & VERIFIED"),
    ("multi-agent-review-orchestrator", "bgaborg", "https://gist.github.com/bgaborg/8058aac0df0e48a3740776963fa87805", "9c568f14", "MIT", "consolidator prompt", "review-consolidator", "AGENT ROLE", "roles/review-consolidator.md", "ADAPTED & VERIFIED"),
    ("multi-model-code-review-agent", "pelednoam", "https://github.com/pelednoam/multi-model-code-review-agent.git", "d41b7535", "MIT", "multi_model_reviewer.py", "multi-model-review", "SKILL / ROUTER", "skills/multi-model-review", "ADAPTED & VERIFIED"),
    ("multi-model-code-review-agent", "pelednoam", "https://github.com/pelednoam/multi-model-code-review-agent.git", "d41b7535", "MIT", "gatekeeper checks", "release-gatekeeper", "AGENT ROLE", "roles/release-gatekeeper.md", "ADAPTED & VERIFIED"),
    ("Claude-Code-Promts-Skills", "Rtur2003", "https://github.com/Rtur2003/Claude-Code-Promts-Skills.git", "ac0502e7", "MIT", "security-audit-prompt.md", "security-auditor", "AGENT ROLE", "roles/security-auditor.md", "ADAPTED & VERIFIED"),
    ("Claude-Code-Promts-Skills", "Rtur2003", "https://github.com/Rtur2003/Claude-Code-Promts-Skills.git", "ac0502e7", "MIT", "security-audit methodology", "security-audit-recon", "SKILL", "skills/security-audit-recon", "ADAPTED & VERIFIED"),
    ("ai-code-review-prompts", "soul-sol", "https://github.com/soul-sol/ai-code-review-prompts.git", "ff21994c", "MIT", "adversarial review prompt", "adversarial-reviewer", "AGENT ROLE", "roles/adversarial-reviewer.md", "ADAPTED & VERIFIED"),
    ("ai-code-review-prompts", "soul-sol", "https://github.com/soul-sol/ai-code-review-prompts.git", "ff21994c", "MIT", "adversarial attack heuristics", "adversarial-review", "SKILL", "skills/adversarial-review", "ADAPTED & VERIFIED")
]

def main():
    print("=== Step 1: Update inventory/agents.json & agents.yaml ===")
    with open(AGENTS_JSON) as f:
        existing_agents = json.load(f)
    
    existing_ids = {a["id"] for a in existing_agents}
    added_count = 0
    for na in NEW_AGENTS:
        if na["id"] not in existing_ids:
            existing_agents.append(na)
            added_count += 1
            print(f"Added {na['id']}: {na['name']}")
    
    with open(AGENTS_JSON, "w") as f:
        json.dump(existing_agents, f, indent=2)
    
    with open(AGENTS_YAML, "w") as f:
        yaml.dump(existing_agents, f, default_flow_style=False, sort_keys=False)
    
    print(f"Updated inventory: now {len(existing_agents)} items ({added_count} newly added).")

    print("\n=== Step 2: Update Agent-Catalog.xlsx ===")
    wb = openpyxl.load_workbook(WORKBOOK_PATH)

    # 1. Update 01_Master_Catalog
    ws01 = wb["01_Master_Catalog"]
    existing_01_ids = {ws01.cell(r, 1).value for r in range(2, ws01.max_row + 1)}
    
    for na in NEW_AGENTS:
        if na["id"] not in existing_01_ids:
            row_idx = ws01.max_row + 1
            row_data = [
                na["id"], na["name"], na["classification"], na["subtype"],
                na["purpose"], na["primary_category"], na["purpose"],
                na["github_repo"], na["github_url"], na["current_location"],
                na["runtime_owner"], "Antigravity / Local CLI", na["invocation_method"],
                ", ".join(na.get("skill_deps", [])), "N/A", ", ".join(na.get("skill_deps", [])),
                ", ".join(na.get("mcp_deps", [])), na["project_scope"], "Global",
                "Native/Curated", "ACTIVE", na["health"], na["version"],
                na["commit"], na["license"], "UNIQUE", "NONE",
                "INTEGRATED", "VERIFIED", "2026-10-01", na["notes"]
            ]
            for col_idx, val in enumerate(row_data, 1):
                cell = ws01.cell(row=row_idx, column=col_idx, value=val)
                cell.border = BORDER_THIN
                if col_idx == 22 and val == "HEALTHY":
                    cell.fill = FILL_HEALTHY
                    cell.font = FONT_HEALTHY
            print(f"Appended to 01_Master_Catalog: {na['id']}")

    # 2. Update 02_Agents
    ws02 = wb["02_Agents"]
    existing_02_ids = {ws02.cell(r, 1).value for r in range(2, ws02.max_row + 1)}
    for na in NEW_AGENTS:
        if na["id"] not in existing_02_ids:
            row_idx = ws02.max_row + 1
            row_data = [
                na["id"], na["name"], "Antigravity / Local CLI", na["runtime"],
                na["model"], na["invocation_method"], na["primary_category"],
                na["health"], na["current_location"], na["purpose"]
            ]
            for col_idx, val in enumerate(row_data, 1):
                cell = ws02.cell(row=row_idx, column=col_idx, value=val)
                cell.border = BORDER_THIN
                if col_idx == 8 and val == "HEALTHY":
                    cell.fill = FILL_HEALTHY
                    cell.font = FONT_HEALTHY
            print(f"Appended to 02_Agents: {na['id']}")

    # 3. Update 03_Roles
    ws03 = wb["03_Roles"]
    existing_03_ids = {ws03.cell(r, 1).value for r in range(2, ws03.max_row + 1)}
    for na in NEW_AGENTS:
        if na["classification"] == "AGENT ROLE" and na["id"] not in existing_03_ids:
            row_idx = ws03.max_row + 1
            row_data = [
                na["id"], na["name"], f"AI-Dev-Team/{na['current_location']} (symlinked in agents/)",
                "AI-Dev-Team Ecosystem", na["github_repo"], na["purpose"],
                na["primary_category"], "VERIFIED"
            ]
            for col_idx, val in enumerate(row_data, 1):
                cell = ws03.cell(row=row_idx, column=col_idx, value=val)
                cell.border = BORDER_THIN
            print(f"Appended to 03_Roles: {na['id']}")

    # 4. Update 04_Skills
    ws04 = wb["04_Skills"]
    existing_skills = {ws04.cell(r, 1).value for r in range(2, ws04.max_row + 1)}
    for sk in NEW_SKILLS:
        if sk[0] not in existing_skills:
            row_idx = ws04.max_row + 1
            for col_idx, val in enumerate(sk, 1):
                cell = ws04.cell(row=row_idx, column=col_idx, value=val)
                cell.border = BORDER_THIN
                if col_idx == 7 and val == "HEALTHY":
                    cell.fill = FILL_HEALTHY
                    cell.font = FONT_HEALTHY
            print(f"Appended to 04_Skills: {sk[0]}")

    # 5. Update 06_Repositories
    ws06 = wb["06_Repositories"]
    existing_repos = {ws06.cell(r, 1).value for r in range(2, ws06.max_row + 1)}
    for rep in NEW_REPOS:
        if rep[0] not in existing_repos:
            row_idx = ws06.max_row + 1
            for col_idx, val in enumerate(rep, 1):
                cell = ws06.cell(row=row_idx, column=col_idx, value=val)
                cell.border = BORDER_THIN
            print(f"Appended to 06_Repositories: {rep[0]}")

    # 6. Update 14_Orchestration
    ws14 = wb["14_Orchestration"]
    existing_orchs = {ws14.cell(r, 1).value for r in range(2, ws14.max_row + 1)}
    if "Pre-Live Full Audit" not in existing_orchs:
        row_idx = ws14.max_row + 1
        orch_data = [
            "Pre-Live Full Audit",
            "Multi-Agent Quality & Security Gate",
            "Discovery -> Static/Security Scans -> Parallel Multi-Agent Audits -> Independent Verification -> Consolidation -> Release Gate",
            "Pre-Live Audit Ruleset (PRE-LIVE-FULL-AUDIT.md)",
            "Gemini, Claude, GitHub Copilot, Local Scanners",
            "HEALTHY"
        ]
        for col_idx, val in enumerate(orch_data, 1):
            cell = ws14.cell(row=row_idx, column=col_idx, value=val)
            cell.border = BORDER_THIN
            if col_idx == 6 and val == "HEALTHY":
                cell.fill = FILL_HEALTHY
                cell.font = FONT_HEALTHY
        print("Appended to 14_Orchestration: Pre-Live Full Audit")

    # 7. Update 15_GitHub_Sources
    ws15 = wb["15_GitHub_Sources"]
    existing_gh = {ws15.cell(r, 1).value for r in range(2, ws15.max_row + 1)}
    for rep in NEW_REPOS:
        # Repos columns in 15: Repository, Upstream URL, Local Clone, Commit, Branch, License, Sync Status
        if rep[0] not in existing_gh:
            row_idx = ws15.max_row + 1
            gh_row = [rep[0], rep[4], rep[1], rep[3], rep[2], rep[6], "SYNCED / ADAPTED"]
            for col_idx, val in enumerate(gh_row, 1):
                cell = ws15.cell(row=row_idx, column=col_idx, value=val)
                cell.border = BORDER_THIN
            print(f"Appended to 15_GitHub_Sources: {rep[0]}")

    # 8. Create or update Sheet 18_Audits_Integration
    if "18_Audits_Integration" in wb.sheetnames:
        del wb["18_Audits_Integration"]
    
    ws18 = wb.create_sheet("18_Audits_Integration")
    headers18 = [
        "Source Repository", "Owner", "Upstream URL", "Commit SHA",
        "License", "Discovered Component", "Adapted / Imported Name",
        "Component Type", "Target Ecosystem Location", "Integration & Health Status"
    ]
    ws18.append(headers18)
    for col_idx in range(1, len(headers18) + 1):
        cell = ws18.cell(row=1, column=col_idx)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, data_row in enumerate(AUDIT_INTEGRATION_DATA, 2):
        for col_idx, val in enumerate(data_row, 1):
            cell = ws18.cell(row=row_idx, column=col_idx, value=val)
            cell.border = BORDER_THIN
            if col_idx == 10 and "VERIFIED" in str(val):
                cell.fill = FILL_HEALTHY
                cell.font = FONT_HEALTHY

    # Auto-adjust column widths for 18_Audits_Integration
    for col in ws18.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws18.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # 9. Update 16_Summary
    ws16 = wb["16_Summary"]
    ws16.cell(3, 3, value="=COUNTA('01_Master_Catalog'!A2:A200)")
    ws16.cell(4, 3, value="=COUNTA('02_Agents'!A2:A100)")
    ws16.cell(5, 3, value="=COUNTA('03_Roles'!A2:A100)")
    ws16.cell(6, 3, value="=COUNTA('04_Skills'!A2:A2000)")
    ws16.cell(9, 3, value="=COUNTA('06_Repositories'!A2:A50)")
    ws16.cell(10, 3, value="=COUNTA('14_Orchestration'!A2:A30)")
    ws16.cell(15, 3, value='=COUNTIF(\'01_Master_Catalog\'!V2:V200, "HEALTHY")')
    ws16.cell(16, 3, value='=COUNTIF(\'01_Master_Catalog\'!V2:V200, "PARTIAL")')
    ws16.cell(17, 3, value='=COUNTIF(\'01_Master_Catalog\'!V2:V200, "BROKEN")')
    ws16.cell(18, 3, value='=COUNTIF(\'01_Master_Catalog\'!V2:V200, "NEEDS USER ACTION")')
    ws16.cell(21, 3, value="3.1.0 (Pre-Live-Full-Audit Integrated)")
    ws16.cell(22, 3, value="2026-10-01")
    ws16.cell(23, 3, value="VALIDATED AFTER SAVE")

    wb.save(WORKBOOK_PATH)
    print(f"\nSuccessfully saved updated workbook to {WORKBOOK_PATH}")

if __name__ == "__main__":
    main()

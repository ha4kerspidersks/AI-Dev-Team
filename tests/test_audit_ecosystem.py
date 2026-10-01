#!/usr/bin/env python3
"""
Comprehensive Automated Test Suite for PRE-LIVE-FULL-AUDIT Integration
Validates:
1. Symlink integrity (agents/, roles/, skills/) -> 0 broken symlinks
2. Pre-Live-Full-Audit Python module imports & classes
3. Technology stack detection against test fixtures
4. Finding schema compliance
5. Finding verification (rejection of non-existent files/out-of-bounds lines)
6. Finding consolidation & Release Gate evaluation
7. Reporter generation (13 mandatory deliverables)
8. Master Agent-Catalog.xlsx integrity & schema match
9. PreLiveFullAudit end-to-end execution
"""

import os
import sys
import json
import shutil
import tempfile
import unittest

ROOT = "/Users/subhajkar/Developer/AI-Dev-Team"
sys.path.insert(0, os.path.join(ROOT, "orchestrators/pre-live-full-audit"))

from finding_schema import Finding, Severity, Status, Category, ReleaseGateResult
from tech_detector import TechDetector, ProjectTechProfile
from verifier import FindingVerifier
from consolidator import FindingConsolidator
from reporter import AuditReporter
from orchestrator import PreLiveFullAudit
import openpyxl

class TestAuditEcosystem(unittest.TestCase):

    def test_01_symlink_integrity(self):
        """Verify 0 broken symlinks exist across agents, roles, and skills."""
        broken_links = []
        directories = [
            os.path.join(ROOT, "agents"),
            os.path.join(ROOT, "roles"),
            os.path.join(ROOT, "skills")
        ]
        
        for d in directories:
            if not os.path.exists(d):
                continue
            for root, dirs, files in os.walk(d):
                for item in files + dirs:
                    p = os.path.join(root, item)
                    if os.path.islink(p):
                        target = os.readlink(p)
                        if not os.path.exists(p):
                            broken_links.append((p, target))

        self.assertEqual(len(broken_links), 0, f"Found broken symlinks: {broken_links}")
        print("PASS: 0 broken symlinks found across agents/, roles/, skills/.")

    def test_02_tech_detector(self):
        """Verify tech stack detector identifies languages, frameworks, DBs, and CI."""
        with tempfile.TemporaryDirectory() as tmpdir:
            with open(os.path.join(tmpdir, "package.json"), "w") as f:
                json.dump({"dependencies": {"react": "^18.2.0", "express": "^4.18.2", "pg": "^8.11.0"}}, f)
            with open(os.path.join(tmpdir, "Dockerfile"), "w") as f:
                f.write("FROM node:18-alpine\n")
            os.makedirs(os.path.join(tmpdir, ".github/workflows"), exist_ok=True)
            with open(os.path.join(tmpdir, ".github/workflows/ci.yml"), "w") as f:
                f.write("name: CI\n")

            detector = TechDetector(tmpdir)
            profile = detector.detect()

            self.assertTrue("JavaScript" in profile.languages or "JavaScript/TypeScript" in profile.languages)
            self.assertIn("React", profile.frameworks)
            self.assertIn("Express.js", profile.frameworks)
            self.assertIn("PostgreSQL", profile.databases)
            self.assertTrue(profile.has_docker)
            self.assertIn("GitHub Actions", profile.ci_cd)
            print("PASS: TechDetector correctly profiled multi-tier stack.")

    def test_03_finding_schema_and_verification(self):
        """Verify Finding schema and deterministic verification engine."""
        with tempfile.TemporaryDirectory() as tmpdir:
            test_file = os.path.join(tmpdir, "auth.py")
            with open(test_file, "w") as f:
                f.write("import os\nSECRET_KEY = 'hardcoded_jwt_secret_token_12345'\n")

            # 1. Valid finding with true cited evidence
            finding_valid = Finding(
                id="FIND-001",
                category=Category.SECRETS.value,
                severity=Severity.HIGH.value,
                confidence=0.95,
                title="Hardcoded Secret Key Detected",
                description="Found raw secret key embedded in source.",
                file="auth.py",
                line=2,
                component="auth_service",
                failure_scenario="Attacker extracts JWT secret from git history.",
                evidence="SECRET_KEY = 'hardcoded_jwt_secret_token_12345'",
                reproduction="Inspect line 2 of auth.py",
                impact="Potential unauthorized token generation",
                recommendation="Move secret to environment variable",
                source_agent="security-auditor",
                source_repository="Rtur2003/Claude-Code-Promts-Skills"
            )

            # 2. Hallucinated finding targeting non-existent line/file
            finding_fake = Finding(
                id="FIND-002",
                category=Category.SECURITY.value,
                severity=Severity.CRITICAL.value,
                confidence=0.3,
                title="Hallucinated SQL Injection",
                description="SQL injection in missing file.",
                file="non_existent.py",
                line=999,
                component="db",
                failure_scenario="Fake exploit",
                evidence="db.query('SELECT * FROM users WHERE id = ' + user_input)",
                reproduction="None",
                impact="High",
                recommendation="Use parameterized queries",
                source_agent="hallucinating-agent",
                source_repository="fake-source"
            )

            verifier = FindingVerifier(tmpdir)
            v_res1 = verifier.verify_finding(finding_valid)
            v_res2 = verifier.verify_finding(finding_fake)

            self.assertEqual(v_res1.status, Status.VERIFIED.value)
            self.assertTrue(v_res1.verified)

            self.assertEqual(v_res2.status, Status.FALSE_POSITIVE.value)
            self.assertFalse(v_res2.verified)
            print("PASS: FindingVerifier confirmed genuine finding and rejected hallucination.")

    def test_04_consolidation_and_release_gating(self):
        """Verify finding deduplication, root cause correlation, and release gate decision."""
        f1 = Finding(
            id="FIND-001",
            category=Category.SECURITY.value,
            severity=Severity.CRITICAL.value,
            confidence=0.99,
            title="Exposed Database Password",
            description="Exposed database credentials in config.",
            file="config.py",
            line=5,
            component="database",
            failure_scenario="Direct DB takeover",
            evidence="DB_PASS = 'admin123'",
            reproduction="Check config.py:5",
            impact="Database breach",
            recommendation="Use vault",
            source_agent="security-auditor",
            source_repository="claude-code-agents",
            verified=True,
            status=Status.VERIFIED.value
        )
        f2 = Finding(
            id="FIND-002",
            category=Category.SECURITY.value,
            severity=Severity.HIGH.value,
            confidence=0.95,
            title="Exposed Database Password",
            description="Duplicate report of DB password in config.",
            file="config.py",
            line=5,
            component="database",
            failure_scenario="Direct DB takeover",
            evidence="DB_PASS = 'admin123'",
            reproduction="Check config.py:5",
            impact="Database breach",
            recommendation="Use vault",
            source_agent="adversarial-reviewer",
            source_repository="ai-code-review-prompts",
            verified=True,
            status=Status.VERIFIED.value
        )

        consolidator = FindingConsolidator(project_path="/tmp/test_project")
        merged = consolidator.consolidate([f1, f2])
        self.assertEqual(len(merged), 1, "Duplicate finding should be merged")
        self.assertEqual(merged[0].severity, Severity.CRITICAL.value, "Highest severity should be preserved")

        gate_decision = consolidator.evaluate_release_gate(merged, test_failures=0, build_failures=0)
        self.assertEqual(gate_decision.status, "BLOCKED")
        self.assertIn("CRITICAL", gate_decision.blocking_reasons[0])
        print("PASS: FindingConsolidator deduplicated findings and correctly BLOCKED release gate.")

    def test_05_reporter_deliverables(self):
        """Verify generation of all 13 required audit deliverables."""
        with tempfile.TemporaryDirectory() as tmpdir:
            out_dir = os.path.join(tmpdir, "audit_output")
            profile = TechDetector(tmpdir).detect()
            reporter = AuditReporter(out_dir, profile)

            dummy_finding = Finding(
                id="FIND-001",
                category=Category.CORRECTNESS.value,
                severity=Severity.LOW.value,
                confidence=0.7,
                title="Missing type annotations",
                description="Function lacks type annotations.",
                file="main.py",
                line=1,
                component="core",
                failure_scenario="Minor typing ambiguity",
                evidence="def run(a, b):",
                reproduction="Inspect main.py:1",
                impact="Low maintainability risk",
                recommendation="Add types",
                source_agent="code-auditor",
                source_repository="claude-code-agents",
                verified=True,
                status=Status.VERIFIED.value
            )

            consolidator = FindingConsolidator(tmpdir)
            gate = consolidator.evaluate_release_gate([dummy_finding], test_failures=0, build_failures=0)

            gen_files = reporter.write_all_reports([dummy_finding], gate)
            expected_files = [
                "AUDIT_REPORT.md", "SECURITY_REPORT.md", "QA_REPORT.md",
                "CODE_REVIEW_REPORT.md", "PERFORMANCE_REPORT.md", "DEPENDENCY_REPORT.md",
                "BROWSER_REPORT.md", "ACCESSIBILITY_REPORT.md", "API_REPORT.md",
                "INFRASTRUCTURE_REPORT.md", "FINDINGS.json", "RELEASE_GATE.json",
                "REMEDIATION_PLAN.md"
            ]

            for ef in expected_files:
                p = os.path.join(out_dir, ef)
                self.assertTrue(os.path.exists(p), f"Missing expected deliverable: {ef}")
                self.assertGreater(os.path.getsize(p), 0, f"Deliverable {ef} is empty")

            print(f"PASS: All {len(expected_files)} deliverables generated and verified.")

    def test_06_master_catalog_excel_integrity(self):
        """Verify Agent-Catalog.xlsx contains all 18 sheets and matches inventory."""
        catalog_path = os.path.join(ROOT, "Agent-Catalog.xlsx")
        wb = openpyxl.load_workbook(catalog_path, data_only=True)
        
        expected_sheets = [
            "01_Master_Catalog", "02_Agents", "03_Roles", "04_Skills", "05_MCP",
            "06_Repositories", "07_Capabilities", "08_Dependencies", "09_Duplicates",
            "10_Health", "11_Migration", "12_Project_Local", "13_Native_Antigravity",
            "14_Orchestration", "15_GitHub_Sources", "16_Summary", "17_Extensions",
            "18_Audits_Integration"
        ]
        
        for es in expected_sheets:
            self.assertIn(es, wb.sheetnames, f"Missing expected sheet: {es}")

        ws01 = wb["01_Master_Catalog"]
        self.assertGreaterEqual(ws01.max_row, 113)

        ws18 = wb["18_Audits_Integration"]
        self.assertGreaterEqual(ws18.max_row, 22)

        with open(os.path.join(ROOT, "inventory/agents.json")) as f:
            agents_data = json.load(f)
        self.assertEqual(len(agents_data), 112)

        print("PASS: Agent-Catalog.xlsx & inventory/agents.json validated (18 sheets, 112 cataloged components).")

    def test_07_pre_live_full_audit_orchestrator_pipeline(self):
        """Verify end-to-end PreLiveFullAudit pipeline execution."""
        with tempfile.TemporaryDirectory() as tmpdir:
            with open(os.path.join(tmpdir, "app.py"), "w") as f:
                f.write("import os\napi_token = 'ghp_samplefaketokenforunittesting1234567890'\nprint('running')\n")
            with open(os.path.join(tmpdir, "requirements.txt"), "w") as f:
                f.write("flask==2.0.1\nrequests==2.28.1\n")

            out_dir = os.path.join(tmpdir, "reports")
            auditor = PreLiveFullAudit(project_path=tmpdir, output_dir=out_dir, read_only=True)
            res = auditor.run(mode="full")

            self.assertIn("profile", res)
            self.assertIn("findings", res)
            self.assertIn("gate", res)
            self.assertIn("reports", res)
            self.assertEqual(len(res["reports"]), 13)
            self.assertIn(res["gate"]["status"], ["BLOCKED", "READY_FOR_REVIEW", "PASS"])
            print(f"PASS: End-to-end PreLiveFullAudit pipeline executed cleanly (Gate: {res['gate']['status']}).")

if __name__ == "__main__":
    unittest.main()
